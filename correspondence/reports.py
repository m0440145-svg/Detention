import io
from collections import Counter
from datetime import timedelta
from django.utils import timezone
from . import services
from .models import *

REPORTS={'register':'السجل الرسمي','weekly-overdue':'المتأخرات الأسبوعية','monthly-unit':'أداء الوحدات','retention':'مرشحو الاستبقاء','permissions':'مراجعة الصلاحيات','secret-views':'الاطلاع على المحمي'}

def metrics(user):
    rows=list(services.visible(user));now=timezone.now();closed=[m for m in rows if m.completed_at]
    ont=[m for m in closed if m.due_at and m.completed_at<=m.due_at]
    refs=[services.work_hours_between(m.created_at,m.first_referred_at) for m in rows if m.first_referred_at]
    cycles=[services.work_hours_between(m.created_at,m.completed_at) for m in closed]
    buckets=Counter();volume=Counter();supervisory=[m for m in closed if m.party.supervisory]
    for m in rows:
        volume[(m.created_at.strftime('%Y-%m'),m.kind,m.party.name,m.classification.name)]+=1
        if m.is_overdue:
            days=(now-m.due_at).total_seconds()/86400
            buckets['0–3' if days<=3 else '4–7' if days<=7 else '8–15' if days<=15 else '>15']+=1
    events=MailEvent.objects.filter(mail__in=services.visible(user))
    return {'total':len(rows),'completed':len(closed),'sla_percent':round(100*len(ont)/len(closed),1) if closed else None,'sla_target':90,'first_referral_work_hours':round(sum(refs)/len(refs),1) if refs else None,'completion_work_hours':round(sum(cycles)/len(cycles),1) if cycles else None,'late_buckets':dict(buckets),'monthly_volumes':[{'month':k[0],'kind':k[1],'party':k[2],'classification':k[3],'count':v} for k,v in volume.items() if not any(m.secrecy in ['secret','restricted'] and m.party.name==k[2] for m in rows)],'supervisory_sla_percent':round(100*sum(m.completed_at<=m.due_at for m in supervisory)/len(supervisory),1) if supervisory else None,'approval_cycles':events.filter(action='إجراء: submit').count(),'rejections':events.filter(action='إجراء: reject').count()}

def report_rows(user,kind):
    from django.core.exceptions import PermissionDenied,ValidationError
    if kind not in REPORTS:raise ValidationError('تقرير غير معروف.')
    if kind=='permissions':
        if user.role!='executive':raise PermissionDenied('مراجعة الصلاحيات للمدير التنفيذي.')
        from core.models import User
        return ['الموظف','الدور','نشط','الوحدات','آخر مزامنة'],[[str(u),u.get_role_display(),u.is_active,', '.join(u.units.values_list('name',flat=True)),str(getattr(getattr(u,'externalidentity',None),'synced_at',''))] for u in User.objects.all()]
    qs=services.visible(user).exclude(secrecy__in=['restricted','secret'])
    if kind=='secret-views':
        if user.role!='executive' and not services.has_access(user,'compliance'):raise PermissionDenied()
        events=MailEvent.objects.filter(mail__in=services.visible(user).filter(secrecy__in=['restricted','secret']),action__icontains='اطلاع')
        return ['مرجع','المطلع','الوقت'],[[e.mail.code,str(e.actor),e.created_at.isoformat()] for e in events]
    if kind=='weekly-overdue':qs=qs.filter(due_at__lt=timezone.now(),paused_at__isnull=True).exclude(status__in=services.TERMINAL)
    if kind=='retention':qs=qs.filter(status='archived',retain_until__lte=timezone.localdate()).exclude(legalhold__active=True)
    if kind=='monthly-unit':
        today=timezone.localdate();qs=qs.filter(created_at__year=today.year,created_at__month=today.month)
        counts=Counter((m.unit.name,m.status) for m in qs)
        return ['الوحدة','الحالة','العدد'],[[a,b,n] for (a,b),n in counts.items()]
    return ['رقم القيد','النوع','الموضوع','الجهة','الوحدة','الحالة','هجري','ميلادي','الاستحقاق'],[[m.code,m.get_kind_display(),m.subject,m.party.name,m.unit.name,m.get_status_display(),m.hijri_date,m.created_at.isoformat(),str(m.due_at or '')] for m in qs]

def xlsx_bytes(headers,rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    workbook=Workbook();sheet=workbook.active;sheet.title='تقرير';sheet.sheet_view.rightToLeft=True
    sheet.append(headers)
    for row in rows:
        # Never interpret user content as Excel formulas.
        sheet.append(["'"+x if isinstance(x,str) and x.startswith(('=','+','-','@')) else x for x in row])
    for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='126854')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=24
    output=io.BytesIO();workbook.save(output);return output.getvalue()

def report_pdf(headers,rows):
    from weasyprint import HTML
    from django.utils.html import escape
    from .documents import local_fetcher
    table='<table><tr>'+''.join('<th>'+escape(x)+'</th>' for x in headers)+'</tr>'+''.join('<tr>'+''.join('<td>'+escape(str(x))+'</td>' for x in row)+'</tr>' for row in rows)+'</table>'
    return HTML(string='<html dir="rtl" lang="ar"><meta charset="utf-8"><style>@page{size:A4 landscape}body{font-family:DejaVu Sans;font-size:9pt}td,th{border:1px solid #aaa;padding:4px}table{border-collapse:collapse}</style>'+table+'</html>',url_fetcher=local_fetcher).write_pdf(pdf_variant='pdf/a-3b')

def run_schedules():
    from .integrations import enqueue
    for schedule in ReportSchedule.objects.filter(active=True,next_run__lte=timezone.now()).select_related('user'):
        if not schedule.user.is_active:schedule.active=False;schedule.save();continue
        try:
            headers,rows=report_rows(schedule.user,schedule.report)
            # Each run re-evaluates the recipient's current access. No protected attachments in mail.
            for config in Integration.objects.filter(active=True,kind='email'):
                enqueue(config,'report.scheduled',{'to':schedule.user.email,'report':schedule.report,'headers':headers,'rows':rows},f'report:{schedule.pk}:{config.pk}:{schedule.next_run.isoformat()}')
            schedule.last_run=timezone.now()
            if schedule.frequency=='monthly':
                import calendar
                date=schedule.last_run;month=date.month%12+1;year=date.year+(date.month==12)
                schedule.next_run=date.replace(year=year,month=month,day=min(date.day,calendar.monthrange(year,month)[1]))
            else:schedule.next_run=schedule.last_run+timedelta(days=1 if schedule.frequency=='daily' else 7)
            schedule.save()
        except Exception as exc:
            from django.core.exceptions import PermissionDenied,ValidationError
            if isinstance(exc,(PermissionDenied,ValidationError)):schedule.active=False;schedule.save()
            else:raise
