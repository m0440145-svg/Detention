import csv, io
from datetime import timedelta
from django.contrib import messages
from django.contrib.auth import authenticate,login,logout
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError,PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Q,Avg,Count
from django.http import HttpResponse, FileResponse, Http404
from django.shortcuts import render,redirect,get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from .models import *
from .forms import *
from .policy import *
from . import services

def signin(request):
    if request.user.is_authenticated: return redirect('dashboard')
    error=''
    if request.method=='POST':
        from .auth import authenticate_attempt,LoginRateLimited
        try:
            user=authenticate_attempt(request,request.POST.get('username',''),request.POST.get('password',''))
            if user:
                from correspondence.security import check_login_factor
                if not check_login_factor(request,user,request.POST.get("otp","")):
                    if request.session.get("mfa_enroll_user"): return redirect("/security/enroll/")
                    return render(request,"core/login.html",{"error":"رمز المصادقة المتعددة غير صحيح."})
                login(request,user)
                request.session["mfa_verified"]=True
                services.audit(None,user,'تسجيل دخول ناجح')
                return redirect('dashboard')
            error='بيانات الدخول غير صحيحة أو الحساب غير نشط.'
        except LoginRateLimited as exc: error=str(exc)
    return render(request,'core/login.html',{'error':error})
@login_required
@require_POST
def signout(request):
    logout(request); return redirect('login')

def filter_tasks(request,qs):
    query=request.GET.get('q','').strip()
    if query:
        term=Q(title__icontains=query)|Q(description__icontains=query)|Q(owner__first_name__icontains=query)|Q(owner__last_name__icontains=query)|Q(unit__name__icontains=query)|Q(decision__number__icontains=query)|Q(decision__text__icontains=query)|Q(decision__meeting__name__icontains=query)
        code=query.upper().removeprefix('TSK-')
        if code.isdigit(): term|=Q(pk=int(code))
        qs=qs.filter(term)
    for field in ['unit','owner','status','priority','source','project','decision','decision__meeting']:
        value=request.GET.get(field)
        if value:
            if field in ['unit','owner','decision','decision__meeting'] and not value.isdigit(): continue
            qs=qs.filter(**{field:value})
    for key,lookup in [('from','start_date__gte'),('to','due_date__lte')]:
        if request.GET.get(key):
            try: date=__import__('datetime').date.fromisoformat(request.GET[key]); qs=qs.filter(**{lookup:date})
            except ValueError: pass
    for key,lookup in [('progress_min','progress__gte'),('progress_max','progress__lte')]:
        if request.GET.get(key,'').isdigit(): qs=qs.filter(**{lookup:int(request.GET[key])})
    return qs.distinct()
def kpis(qs):
    tasks=list(qs); total=len(tasks); closed=[t for t in tasks if t.status==Status.CLOSED]; open_tasks=[t for t in tasks if t.status not in services.TERMINAL]
    overdue=[t for t in open_tasks if t.is_overdue]; blocked=[t for t in open_tasks if t.is_blocked]
    ontime=sum(1 for t in closed if t.submitted_at and timezone.localtime(t.submitted_at).date()<=t.due_date)
    durations=[(timezone.localtime(t.closed_at).date()-t.start_date).days for t in closed if t.closed_at]
    return {'total':total,'open':len(open_tasks),'active':sum(t.status==Status.ACTIVE for t in tasks),'closed':len(closed),'overdue':len(overdue),'blocked':len(blocked),'critical':sum(t.priority==Priority.CRITICAL and t.status not in services.TERMINAL for t in tasks),'approval':sum(t.status==Status.APPROVAL for t in tasks),'completion':round(len(closed)/total*100) if total else 0,'ontime':round(ontime/len(closed)*100) if closed else None,'mean_days':round(sum(durations)/len(durations),1) if durations else None,'mean_late':round(sum(-t.days_remaining for t in overdue)/len(overdue),1) if overdue else 0,'late_rate':round(len(overdue)/len(open_tasks)*100) if open_tasks else 0,'blocked_rate':round(len(blocked)/len(open_tasks)*100) if open_tasks else 0,'closed_month':sum(t.closed_at and timezone.localtime(t.closed_at).date().replace(day=1)==timezone.localdate().replace(day=1) for t in closed)}
def choices(user):
    visible=tasks_for(user)
    return {'units':Unit.objects.filter(pk__in=visible.values('unit_id')),'employees':User.objects.filter(pk__in=visible.values('owner_id')),'statuses':Status.choices,'priorities':Priority.choices,'sources':Source.choices}
@login_required
def dashboard(request):
    if request.user.role==Role.BOARD: return redirect('decisions')
    if request.user.role==Role.ADMIN:
        return render(request,'core/admin_dashboard.html',{'users_count':User.objects.count(),'active_users':User.objects.filter(is_active=True).count(),'units_count':Unit.objects.count(),'active_units':Unit.objects.filter(active=True).count()})
    qs=tasks_for(request.user); today=timezone.localdate()
    rows=[]
    for unit in Unit.objects.filter(pk__in=qs.values('unit_id')):
        rows.append({'unit':unit,**kpis(qs.filter(unit=unit))})
    employees=[]
    for u in User.objects.filter(pk__in=qs.values('owner_id')):
        employees.append({'user':u,**kpis(qs.filter(owner=u))})
    return render(request,'core/dashboard.html',{'kpi':kpis(qs),'unit_rows':sorted(rows,key=lambda r:r['completion'],reverse=True),'employee_rows':employees,'decision_obstacles':Obstacle.objects.filter(task__in=qs,needs_decision=True,resolved_at__isnull=True).select_related('task','intervention_owner')[:5],'late_week':qs.filter(due_date__lt=today-timedelta(days=7)).exclude(status__in=services.TERMINAL).count(),'today_tasks':qs.filter(due_date=today)[:8],'urgent_tasks':qs.filter(priority=Priority.CRITICAL).exclude(status__in=services.TERMINAL)[:6],'recent_comments':Comment.objects.filter(task__in=qs).select_related('user','task').order_by('-created_at')[:5],'notifications':request.user.notifications.all()[:5]})
@login_required
def task_list(request):
    require_execution_space(request.user)
    qs=filter_tasks(request,tasks_for(request.user))
    scope=request.GET.get('scope')
    if scope=='mine': qs=qs.filter(Q(owner=request.user)|Q(participants=request.user)).distinct()
    if scope=='unit': qs=qs.filter(unit__in=request.user.units.all())
    return render(request,'core/tasks.html',{'page':Paginator(qs,20).get_page(request.GET.get('page')), 'title':'المهام والتكليفات',**choices(request.user)})
@login_required
def task_form(request,pk=None):
    require_execution_space(request.user)
    task=get_object_or_404(tasks_for(request.user),pk=pk) if pk else None
    require_manage(request.user,task)
    initial={}
    decision_id=request.GET.get('decision')
    if decision_id:
        decision=get_object_or_404(decisions_for(request.user),pk=decision_id)
        initial={'decision':decision,'title':decision.text[:250],'description':decision.text,'source':Source.DECISION,'start_date':decision.date,'expected_result':'تنفيذ القرار '+decision.number}
    form=TaskForm(request.POST or None,instance=task,user=request.user,initial=initial)
    if request.method=='POST' and form.is_valid():
        data={k:form.cleaned_data[k] for k in TaskForm.Meta.fields if k not in ['participants','participating_units']}
        try:
            if task: task=services.edit_task(request.user,task,data,form.cleaned_data['participants'],form.cleaned_data['participating_units'],form.cleaned_data['reason'])
            else: task=services.create_task(request.user,data,form.cleaned_data['participants'],form.cleaned_data['participating_units'])
            messages.success(request,'تم حفظ المهمة وإرسال التنبيهات.'); return redirect('task-detail',pk=task.pk)
        except ValidationError as e: form.add_error(None,e)
    return render(request,'core/form.html',{'form':form,'title':'تعديل المهمة' if task else 'إضافة مهمة'})
@login_required
def task_detail(request,pk):
    require_operational(request.user)
    task=get_object_or_404(tasks_for(request.user),pk=pk)
    error=None
    if request.method=='POST':
        action=request.POST.get('action'); require_work(request.user,task)
        try:
            if action=='progress':
                form=ProgressForm(request.POST)
                if not form.is_valid(): raise ValidationError(str(form.errors))
                services.update_progress(request.user,task,**form.cleaned_data)
            elif action=='comment':
                form=CommentForm(request.POST,task=task)
                if not form.is_valid(): raise ValidationError(str(form.errors))
                comment=services.add_comment(request.user,task,**form.cleaned_data)
                if request.FILES.get('file'): services.attach(request.user,task,request.FILES['file'],comment=comment)
            elif action=='obstacle':
                form=ObstacleForm(request.POST,task=task)
                if not form.is_valid(): raise ValidationError(str(form.errors))
                services.report_obstacle(request.user,task,form.cleaned_data)
            elif action=='resolve':
                obstacle=get_object_or_404(task.obstacles,pk=request.POST.get('obstacle'))
                services.resolve_obstacle(request.user,obstacle,request.POST.get('resolution',''))
            elif action=='status': services.change_status(request.user,task,request.POST.get('status'),request.POST.get('comment',''))
            elif action=='attachment':
                upload=request.FILES.get('file')
                if not upload: raise ValidationError('اختر ملفًا.')
                subtask=get_object_or_404(task.subtasks,pk=request.POST['subtask']) if request.POST.get('subtask') else None
                services.attach(request.user,task,upload,subtask=subtask)
            elif action=='escalate': services.escalate(request.user,task,request.POST.get('reason',''))
            elif action=='subtask':
                instance=get_object_or_404(task.subtasks,pk=request.POST['subtask']) if request.POST.get('subtask') else None
                form=SubtaskForm(request.POST,instance=instance)
                if not form.is_valid(): raise ValidationError(str(form.errors))
                services.save_subtask(request.user,task,form.cleaned_data,instance)
            else: raise ValidationError('الإجراء غير معروف.')
            messages.success(request,'تم حفظ الإجراء وتوثيقه في سجل النشاط.')
            return redirect('task-detail',pk=pk)
        except ValidationError as e: error=' — '.join(e.messages)
    task.refresh_from_db()
    statuses=services.TRANSITIONS.get(task.status,set())
    if not can_manage(request.user,task): statuses=statuses-{Status.CLOSED,Status.RETURNED,Status.CANCELLED,Status.PAUSED}
    from correspondence.services import visible as visible_mail
    linked_mail=visible_mail(request.user).filter(referrals__task=task).first()
    return render(request,'core/detail.html',{'linked_mail':linked_mail,'task':task,'error':error,'can_work':can_work(request.user,task),'can_manage':can_manage(request.user,task),'next_statuses':[(s,Catalog.objects.filter(kind='status',key=s).values_list('label',flat=True).first() or Status(s).label) for s in statuses if not Catalog.objects.filter(kind='status',key=s,active=False).exists()],'progress_form':ProgressForm(initial={'progress':task.progress}),'comment_form':CommentForm(task=task),'obstacle_form':ObstacleForm(task=task),'subtask_form':SubtaskForm(),'flags':services.flags(task),'activities':task.activities.select_related('actor'),'comments':task.comments.select_related('user','parent').prefetch_related('mentions','mentioned_units','attachments').order_by('created_at')})
FOLLOW_TABS=[('overdue','المتأخرة'),('blocked','المتعثرة'),('today','تستحق اليوم'),('soon','خلال 3 أيام'),('stale','بدون تحديث'),('waiting','بانتظار رد'),('approval','بانتظار اعتماد'),('critical','الحرجة'),('escalated','المصعدة'),('returned','المعادة'),('risk','مرشحة للتعثر')]
RISK_GROUPS=[('overdue','متأخرة','تجاوزت تاريخ الاستحقاق'),('blocked','متعثرة','عوائق مفتوحة تحتاج معالجة'),('stale','بدون تحديث','لم تسجل تحديثًا خلال المدة المحددة'),('escalated','مصعّدة','أُحيلت للتدخل الإداري')]
def risk_tasks(qs,kind,today=None):
    today=today or timezone.localdate()
    if kind=='overdue': return qs.filter(due_date__lt=today)
    if kind=='blocked': return qs.filter(obstacles__resolved_at__isnull=True,obstacles__isnull=False).distinct()
    if kind=='stale': return qs.filter(updated_at__lt=timezone.now()-timedelta(days=RuleSettings.current().stale_days))
    if kind=='escalated': return qs.filter(escalations__isnull=False).distinct()
    if kind=='today': return qs.filter(due_date=today)
    if kind=='soon': return qs.filter(due_date__range=(today,today+timedelta(days=3)))
    if kind=='critical': return qs.filter(priority=Priority.CRITICAL)
    if kind=='risk': return qs.filter(pk__in=[t.pk for t in qs if 'مرشحة للتعثر' in services.flags(t)])
    return qs.filter(status=kind) if kind in Status.values else qs.none()
@login_required
def followup(request):
    require_execution_space(request.user)
    error=None
    if request.method=='POST':
        try:
            result=services.bulk_followup(request.user,request.POST.getlist('task_ids'),request.POST.get('action',''),request.POST.get('reason',''))
            messages.success(request,f"تمت متابعة {result['processed']} مهمة، وتجاوز {result['skipped']} إجراء مكرر اليوم.")
            return redirect('followup')
        except ValidationError as exc: error=' — '.join(exc.messages)
    qs=filter_tasks(request,tasks_for(request.user)).exclude(status__in=services.TERMINAL)
    tab=request.GET.get('tab','all')
    definitions=RISK_GROUPS if tab=='all' else [(tab,dict(FOLLOW_TABS).get(tab,'المتابعة'),'عرض ضمن نطاق الصلاحيات')]
    groups=[]
    for key,title,description in definitions:
        selected=risk_tasks(qs,key)
        page=Paginator(selected,20).get_page(request.GET.get('page'))
        for task in page: task.bulk_allowed=can_manage(request.user,task)
        groups.append({'key':key,'title':title,'description':description,'count':page.paginator.count,'page':page})
    return render(request,'core/followup.html',{'groups':groups,'page':groups[0]['page'],'tabs':[('all','جميع المخاطر'),*FOLLOW_TABS],'tab':tab,'can_bulk':can_manage(request.user),'error':error,**choices(request.user)})
def decisions_for(user):
    qs=Decision.objects.select_related('meeting','followup_owner','approved_minutes_attachment').prefetch_related('tasks')
    if user.role==Role.ADMIN: return qs.none()
    if user.role in DECISION_READ_ROLES: return qs
    return qs.filter(Q(tasks__in=tasks_for(user))|Q(followup_owner=user)).distinct()
def decision_metrics(user,decision):
    tasks=tasks_for(user).filter(decision=decision); total=tasks.count()
    return {'progress':round(tasks.aggregate(p=Avg('progress'))['p'] or 0),'approved_progress':round(tasks.filter(status=Status.CLOSED).count()/total*100) if total else 0,'total':total,'full_scope':user.role in DECISION_READ_ROLES}
def decision_activities(user,decision):
    qs=decision.activities.select_related('actor')
    # Calculated whole-decision percentages can reveal hidden task performance.
    if user.role not in DECISION_READ_ROLES: qs=qs.filter(Q(new__derived__isnull=True)|Q(new__derived=False))
    return qs
@login_required
def decisions(request):
    require_operational(request.user)
    rows=[]
    for decision in decisions_for(request.user):
        metrics=decision_metrics(request.user,decision)
        decision.visible_progress=metrics['progress']; decision.visible_approved_progress=metrics['approved_progress']
        rows.append(decision)
    return render(request,'core/decisions.html',{'decisions':rows,'manage':request.user.role in GLOBAL_ROLES,'full_scope':request.user.role in DECISION_READ_ROLES,'meetings':Meeting.objects.exclude(governance__secrecy='secret') if request.user.role in GLOBAL_ROLES else Meeting.objects.filter(decisions__in=decisions_for(request.user)).distinct()})
@login_required
def decision_detail(request,pk):
    require_operational(request.user)
    decision=get_object_or_404(decisions_for(request.user),pk=pk)
    error=None
    if request.method=='POST':
        services.require_decision_manage(request.user)
        try:
            upload=request.FILES.get('file')
            if not upload: raise ValidationError('اختر الملف.')
            services.attach_decision(request.user,decision,upload,request.POST.get('reason',''),approved_minutes=request.POST.get('approved_minutes')=='on')
            return redirect('decision-detail',pk=pk)
        except ValidationError as e: error=' — '.join(e.messages)
    return render(request,'core/decision.html',{'decision':decision,'tasks':tasks_for(request.user).filter(decision=decision),'metrics':decision_metrics(request.user,decision),'activities':decision_activities(request.user,decision),'manage':request.user.role in GLOBAL_ROLES,'error':error})
@login_required
def generic_form(request,kind,pk=None):
    definitions={'unit':(Unit,UnitForm,'وحدة تنظيمية'),'employee':(User,EmployeeForm,'موظف'),'meeting':(Meeting,MeetingForm,'اجتماع'),'decision':(Decision,DecisionForm,'قرار اجتماع')}
    if kind not in definitions: raise Http404()
    if kind=='meeting':
        if not pk:
            if request.user.role not in GLOBAL_ROLES:raise PermissionDenied()
            return redirect('meeting-new')
        from meetinghub.models import MeetingRecord
        expanded=MeetingRecord.objects.filter(meeting_id=pk).first()
        if expanded:
            from meetinghub.services import records_for
            get_object_or_404(records_for(request.user,True),pk=expanded.pk)
            return redirect('meeting-edit',pk=expanded.pk)
    if kind in ['unit','employee'] and __import__('correspondence.models',fromlist=['Integration']).Integration.objects.filter(kind='hr',active=True).exists(): raise PermissionDenied('بيانات الموظفين والوحدات تأتي من الموارد البشرية؛ الإدارة اليدوية موقوفة أثناء الربط.')
    if kind in ['unit','employee'] and request.user.role!=Role.ADMIN: raise PermissionDenied()
    if kind in ['meeting','decision'] and request.user.role not in GLOBAL_ROLES: raise PermissionDenied()
    model,klass,label=definitions[kind]
    instance=get_object_or_404(model,pk=pk) if pk else None
    if kind=='employee' and pk: klass=EmployeeEditForm
    form=klass(request.POST or None,request.FILES or None,instance=instance)
    if request.method=='POST' and form.is_valid():
        if kind=='decision':
            try:
                data={key:form.cleaned_data[key] for key in DecisionForm.Meta.fields if key!='minutes_file'}
                item=services.save_decision(request.user,data,instance,form.cleaned_data.get('reason',''),form.cleaned_data.get('minutes_file'))
            except ValidationError as e:
                form.add_error(None,e)
                return render(request,'core/form.html',{'form':form,'title':('تعديل ' if pk else 'إضافة ')+label})
        else: item=form.save()
        if kind=='unit' and item.head: item.head.units.add(item)
        if kind!='decision': services.audit(None,request.user,'حفظ '+label,new={'id':item.pk})
        messages.success(request,'تم الحفظ.'); return redirect('directory' if kind in ['unit','employee'] else 'decisions')
    return render(request,'core/form.html',{'form':form,'title':('تعديل ' if pk else 'إضافة ')+label})
@login_required
def directory(request):
    if request.user.role==Role.BOARD: raise PermissionDenied()
    qs=tasks_for(request.user)
    units=Unit.objects.all() if request.user.role in DIRECTORY_ROLES else request.user.units.all()
    users=User.objects.all() if request.user.role in DIRECTORY_ROLES else User.objects.filter(units__in=units).distinct() if request.user.role==Role.HEAD else User.objects.filter(pk=request.user.pk)
    unit_rows=[{'unit':u,**kpis(qs.filter(unit=u))} for u in units]
    employee_rows=[{'user':u,**kpis(qs.filter(Q(owner=u)|Q(participants=u)).distinct())} for u in users]
    return render(request,'core/directory.html',{'unit_rows':unit_rows,'employee_rows':employee_rows,'admin':request.user.role==Role.ADMIN})
REPORTS=[('units','أداء الوحدات'),('employees','أداء الموظفين'),('overdue','المهام المتأخرة'),('blocked','المهام المتعثرة'),('ontime','الالتزام بالمواعيد'),('closed','المهام المكتملة'),('stale','المهام بدون تحديث'),('decisions','قرارات الاجتماعات'),('execution','تنفيذ القرارات'),('workload','ضغط العمل'),('duration','متوسط زمن الإنجاز'),('escalations','التصعيدات')]
def report_rows(qs,kind,user):
    if kind=='units':
        header=['الوحدة','المهام','المغلقة','المتأخرة','المتعثرة','نسبة الإنجاز']
        rows=[]
        for unit in Unit.objects.filter(pk__in=qs.values('unit_id')):
            k=kpis(qs.filter(unit=unit)); rows.append([unit.name,k['total'],k['closed'],k['overdue'],k['blocked'],k['completion']])
    elif kind in ['employees','workload']:
        header=['الموظف','إجمالي التكليفات','مفتوحة','مغلقة','متأخرة','متعثرة','الإنجاز %']; rows=[]
        for u in User.objects.filter(Q(owned_tasks__in=qs)|Q(participating_tasks__in=qs)).distinct():
            k=kpis(qs.filter(Q(owner=u)|Q(participants=u)).distinct()); rows.append([str(u),k['total'],k['open'],k['closed'],k['overdue'],k['blocked'],k['completion']])
    elif kind in ['decisions','execution']:
        header=['القرار','الاجتماع','النص','المهام المرئية','الإنجاز المرئي %','الاعتماد المرئي %']; rows=[]
        for d in decisions_for(user):
            tasks=qs.filter(decision=d); n=tasks.count()
            rows.append([d.number,d.meeting.name,d.text,n,round(tasks.aggregate(p=Avg('progress'))['p'] or 0),round(tasks.filter(status=Status.CLOSED).count()/n*100) if n else 0])
    elif kind=='escalations':
        header=['المهمة','المستلم','المستوى','السبب','التاريخ']; rows=[[e.task.code,str(e.recipient),e.level,e.reason,e.created_at.strftime('%Y-%m-%d')] for e in Escalation.objects.filter(task__in=qs).select_related('task','recipient')]
    else:
        if kind=='overdue': qs=qs.filter(due_date__lt=timezone.localdate()).exclude(status__in=services.TERMINAL)
        if kind=='blocked': qs=qs.exclude(status__in=services.TERMINAL).filter(obstacles__resolved_at__isnull=True,obstacles__isnull=False).distinct()
        if kind in ['closed','ontime','duration']: qs=qs.filter(status=Status.CLOSED)
        if kind=='stale': qs=qs.exclude(status__in=services.TERMINAL).filter(updated_at__lt=timezone.now()-timedelta(days=RuleSettings.current().stale_days))
        header=['رقم المهمة','العنوان','الوحدة','المسؤول','المصدر','الحالة','الأولوية','الإنجاز %','الاستحقاق','آخر تحديث','تاريخ التسليم','تاريخ الإغلاق','أيام التأخير','مدة التعثر (أيام)','زمن الإنجاز (أيام)']
        rows=[]
        for t in qs:
            intervals=sorted([(o.created_at,o.resolved_at or timezone.now()) for o in t.obstacles.all()]); merged=[]
            for start,end in intervals:
                if merged and start<=merged[-1][1]: merged[-1]=(merged[-1][0],max(end,merged[-1][1]))
                else: merged.append((start,end))
            blocked_days=round(sum((b-a).total_seconds() for a,b in merged)/86400,1)
            rows.append([t.code,t.title,t.unit.name,str(t.owner),t.get_source_display(),t.get_status_display(),t.get_priority_display(),t.progress,str(t.due_date),t.updated_at.strftime('%Y-%m-%d'),t.submitted_at.strftime('%Y-%m-%d') if t.submitted_at else '',t.closed_at.strftime('%Y-%m-%d') if t.closed_at else '',max(0,-t.days_remaining) if t.is_overdue else 0,blocked_days,(timezone.localtime(t.closed_at).date()-t.start_date).days if t.closed_at else ''])
    return header,rows
@login_required
def reports(request):
    require_operational(request.user)
    qs=filter_tasks(request,tasks_for(request.user)); kind=request.GET.get('report','execution' if request.user.role==Role.BOARD else 'units')
    if request.user.role==Role.BOARD and kind not in ['decisions','execution']: raise PermissionDenied('دور المجلس يقرأ تقارير القرارات وتنفيذها فقط.')
    if kind not in dict(REPORTS): kind='units'
    header,rows=report_rows(qs,kind,request.user)
    export=request.GET.get('export')
    if export=='csv':
        response=HttpResponse(content_type='text/csv; charset=utf-8'); response['Content-Disposition']='attachment; filename="report.csv"'; response.write('\ufeff')
        writer=csv.writer(response); writer.writerow(header)
        for row in rows: writer.writerow([safe_cell(v) for v in row])
        return response
    if export=='xlsx':
        book=Workbook(); sheet=book.active; sheet.title='التقرير'; sheet.sheet_view.rightToLeft=True; sheet.append(header)
        for row in rows: sheet.append([safe_cell(v) for v in row])
        from openpyxl.styles import Font,PatternFill
        for c in sheet[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='CF6028')
        for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=28
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        buffer=io.BytesIO(); book.save(buffer)
        response=HttpResponse(buffer.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); response['Content-Disposition']='attachment; filename="report.xlsx"'; return response
    return render(request,'core/reports.html',{'header':header,'rows':rows,'kpi':kpis(qs),'report_types':[item for item in REPORTS if item[0] in ['decisions','execution']] if request.user.role==Role.BOARD else REPORTS,'report':kind,**choices(request.user)})
def safe_cell(value):
    return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value
@login_required
def notifications(request):
    require_execution_space(request.user)
    if request.method=='POST':
        if request.POST.get('id'): request.user.notifications.filter(pk=request.POST['id']).update(read=True)
        else: request.user.notifications.filter(read=False).update(read=True)
        return redirect('notifications')
    return render(request,'core/notifications.html',{'page':Paginator(request.user.notifications.select_related('task'),30).get_page(request.GET.get('page'))})
@login_required
def download(request,pk):
    a=get_object_or_404(Attachment,pk=pk)
    task_id=a.task_id or (a.subtask.task_id if a.subtask_id else None)
    allowed=tasks_for(request.user).filter(pk=task_id).exists() if task_id else decisions_for(request.user).filter(pk=a.decision_id).exists()
    if not allowed: raise Http404()
    try: stream=a.file.open('rb')
    except FileNotFoundError: raise Http404()
    response=FileResponse(stream,as_attachment=True,filename=a.original_name); response['X-Content-Type-Options']='nosniff'; return response
@login_required
def settings(request):
    if request.user.role!=Role.ADMIN: raise PermissionDenied()
    form=RulesForm(request.POST or None,instance=RuleSettings.current())
    if request.method=='POST' and form.is_valid():
        form.save(); services.audit(None,request.user,'تعديل قواعد المتابعة والتصعيد',new=form.cleaned_data); messages.success(request,'تم تحديث القواعد.'); return redirect('settings')
    return render(request,'core/form.html',{'form':form,'title':'قواعد المتابعة والتصعيد'})

@login_required
def catalogs(request):
    if request.user.role!=Role.ADMIN: raise PermissionDenied()
    for kind,values in [('status',Status.choices),('priority',Priority.choices),('task_type',[('تشغيلية','تشغيلية'),('مشروع','مشروع'),('إدارية','إدارية')])]:
        for key,label in values: Catalog.objects.get_or_create(kind=kind,key=key,defaults={'label':label})
    error=None
    if request.method=='POST':
        kind=request.POST.get('kind'); key=request.POST.get('key','').strip()[:100]; label=request.POST.get('label','').strip()[:100]; active=request.POST.get('active')=='on'
        protected={Status.ASSIGNED,Status.ACTIVE,Status.APPROVAL,Status.RETURNED,Status.CLOSED,Status.CANCELLED}
        if kind not in ['status','priority','task_type'] or not key or not label: error='أدخل مفتاحًا واسمًا صحيحين.'
        elif kind=='status' and (key not in Status.values or (key in protected and not active)): error='لا يمكن إضافة حالة خارج سير العمل أو إيقاف حالة أساسية.'
        elif kind=='priority' and (key not in Priority.values or (key==Priority.NORMAL and not active)): error='لا يمكن إيقاف الأولوية الافتراضية أو إضافة مستوى خارج ترتيب الأولويات.'
        else:
            c,_=Catalog.objects.update_or_create(kind=kind,key=key,defaults={'label':label,'active':active})
            services.audit(None,request.user,'تعديل دليل الإعدادات',new={'kind':kind,'key':key,'label':label,'active':active})
            messages.success(request,'تم تحديث الإعدادات مع الحفاظ على سجلات المهام السابقة.'); return redirect('catalogs')
    return render(request,'core/catalogs.html',{'catalogs':Catalog.objects.order_by('kind','pk'),'error':error})
@login_required
def permissions(request):
    if request.user.role!=Role.ADMIN: raise PermissionDenied()
    rows=[['مجلس الإدارة / أمين المجلس','القرارات والمهام المرتبطة وتقارير تنفيذها','لا','لا','قراءة النسب الكلية وسجل تدقيق القرارات'],['مسؤول النظام','المستخدمون والوحدات والإعدادات','لا','لا','إدارة المستخدمين والوحدات والإعدادات فقط'],['المدير التنفيذي / المساعد','جميع الوحدات','نعم','نعم؛ دون الاعتماد الذاتي','الاجتماعات والقرارات'],['رئيس الوحدة','مهام وحدته والمشارك بها','داخل الوحدة القائدة','داخل الوحدة القائدة','متابعة موظفي الوحدة'],['الموظف','مهامه كمسؤول أو مشارك','لا','لا','تحديث التنفيذ والردود'],['العرض فقط','المهام العادية لوحدته','لا','لا','قراءة وتقارير فقط']]
    return render(request,'core/permissions.html',{'rows':rows})


# JSON session authentication uses Django CSRF protection even before login.
from django.http import JsonResponse
from django.middleware.csrf import get_token
from django.views.decorators.csrf import ensure_csrf_cookie,csrf_protect
from django.views.decorators.http import require_GET

@require_GET
@ensure_csrf_cookie
def api_csrf(request):
    return JsonResponse({'csrf_token':get_token(request)})

@require_POST
@csrf_protect
def api_login(request):
    import json
    from .auth import authenticate_attempt,LoginRateLimited
    try:
        data=json.loads(request.body)
        if not isinstance(data,dict): raise ValueError()
        identifier=data.get('username',''); password=data.get('password','')
        if not isinstance(identifier,str) or not isinstance(password,str) or len(identifier)>254 or len(password)>1024: raise ValueError()
    except (ValueError,UnicodeDecodeError):
        return JsonResponse({'detail':'طلب الدخول غير صالح.'},status=400)
    try: user=authenticate_attempt(request,identifier,password)
    except LoginRateLimited as exc: return JsonResponse({'detail':str(exc)},status=429)
    if not user: return JsonResponse({'detail':'بيانات الدخول غير صحيحة أو الحساب غير نشط.'},status=401)
    from correspondence.security import check_login_factor
    if not check_login_factor(request,user,data.get('otp','')): return JsonResponse({'detail':'MFA required','enrollment_url':'/security/enroll/'},status=403)
    login(request,user); request.session['mfa_verified']=True; services.audit(None,user,'تسجيل دخول ناجح')
    return JsonResponse({'id':user.pk,'name':str(user),'role':user.role,'csrf_token':get_token(request)})
